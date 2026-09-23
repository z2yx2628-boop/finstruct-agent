"""Parse Word (.docx) and spreadsheet (.xlsx, .xlsm, .csv) documents.

Neither format has fixed pages, so content is grouped into numbered text
blocks that play the role of pages, as for webpages:

* Word: paragraphs are numbered [P1], [P2] ...; each table row becomes one
  line with cells joined by " | " and is also kept cell-by-cell in `tables`.
* Spreadsheets: every sheet is split into blocks; each block starts with
  "[工作表 名称]" and repeats the header row, and every data line starts
  with its Excel row number, e.g. "[R12] 华菱涟钢 | 300,000".

Legacy Word 97-2003 (.doc) and Excel 97-2003 (.xls) files are read with the
built-in compound-file reader in src.legacy_office (no extra packages). In a
.doc, lines that came from table rows contain " | " and are not numbered as
paragraphs. Older Word 6/95 or Excel 5/95 files and encrypted files raise an
error asking the user to save the file as .docx / .xlsx.
"""
import csv
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from schemas.common import ParsedDocument, ParsedPage
from src.legacy_office import doc_text, xls_sheets
from src.pdf_parser import calculate_sha256

WORD_SUFFIXES = (".docx", ".doc")
SPREADSHEET_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".csv")
BLOCK_CHAR_LIMIT = 1500

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _blocks(lines: list[str], header: list[str] | None = None) -> list[str]:
    """Group lines into blocks of about BLOCK_CHAR_LIMIT characters."""
    header = header or []
    blocks, current = [], list(header)
    for line in lines:
        size = sum(len(item) for item in current)
        if current[len(header):] and size + len(line) > BLOCK_CHAR_LIMIT:
            blocks.append("\n".join(current))
            current = list(header)
        current.append(line)
    if current[len(header):]:
        blocks.append("\n".join(current))
    return blocks


def _paragraph_text(paragraph) -> str:
    parts = []
    for node in paragraph.iter():
        if node.tag == W + "t" and node.text:
            parts.append(node.text)
        elif node.tag == W + "tab":
            parts.append(" ")
        elif node.tag in (W + "br", W + "cr"):
            parts.append("\n")
    return re.sub(r"[ \t]+", " ", "".join(parts)).strip()


def _legacy_doc_lines(path: Path) -> tuple[list[str], list[dict], int]:
    lines, table_rows, paragraph_number = [], [], 0
    for line in doc_text(path).split("\n"):
        if " | " in line:
            lines.append(line)
            table_rows.append(line.split(" | "))
        else:
            paragraph_number += 1
            lines.append(f"[P{paragraph_number}] {line}")
    tables = [{"index": 1, "rows": table_rows}] if table_rows else []
    return lines, tables, paragraph_number


def parse_docx(path: Path) -> ParsedDocument:
    path = Path(path).resolve()
    if path.suffix.lower() == ".doc":
        lines, tables, paragraph_number = _legacy_doc_lines(path)
        return _word_document(path, lines, tables, paragraph_number, legacy=True)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    body = root.find(W + "body")

    lines: list[str] = []
    tables: list[dict] = []
    paragraph_number = 0
    for element in list(body) if body is not None else []:
        if element.tag == W + "p":
            text = _paragraph_text(element)
            if text:
                paragraph_number += 1
                lines.append(f"[P{paragraph_number}] {text}")
        elif element.tag == W + "tbl":
            rows = []
            for row in element.iter(W + "tr"):
                cells = [
                    " ".join(
                        filter(None, (_paragraph_text(p) for p in cell.iter(W + "p")))
                    )
                    for cell in row.findall(W + "tc")
                ]
                if any(cells):
                    rows.append(cells)
            if rows:
                tables.append({"index": len(tables) + 1, "rows": rows})
                lines.extend(" | ".join(row) for row in rows)

    return _word_document(path, lines, tables, paragraph_number)


def _word_document(
    path: Path,
    lines: list[str],
    tables: list[dict],
    paragraph_number: int,
    legacy: bool = False,
) -> ParsedDocument:
    pages = [
        ParsedPage(page=index, text=text)
        for index, text in enumerate(_blocks(lines), start=1)
    ]
    if pages and tables:
        pages[0] = pages[0].model_copy(update={"tables": tables})
    return ParsedDocument(
        document_id=path.stem,
        source_type="word",
        source_name=path.name,
        source_path=str(path),
        sha256=calculate_sha256(path),
        pages=pages,
        metadata={
            "page_count": len(pages),
            "page_unit": "text_block",
            "paragraph_count": paragraph_number,
            "table_count": len(tables),
            "legacy_format": legacy,
            "text_char_count": sum(len(page.text) for page in pages),
            "empty_page_count": 0,
            "needs_ocr": False,
        },
    )


def _sheet_rows(path: Path) -> list[tuple[str, list[list[str]]]]:
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        rows = [[_cell_text(cell) for cell in row] for row in csv.reader(text.splitlines())]
        return [(path.stem, rows)]

    if path.suffix.lower() == ".xls":
        return [
            (name, [[_cell_text(cell) for cell in row] for row in rows])
            for name, rows in xls_sheets(path)
        ]

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        sheets.append((sheet.title, rows))
    workbook.close()
    return sheets


def parse_spreadsheet(path: Path) -> ParsedDocument:
    path = Path(path).resolve()
    pages: list[ParsedPage] = []
    sheet_names = []
    for sheet_name, rows in _sheet_rows(path):
        numbered = [
            (number, [cell for cell in row])
            for number, row in enumerate(rows, start=1)
        ]
        while numbered and not any(numbered[-1][1]):
            numbered.pop()
        numbered = [(n, row) for n, row in numbered if any(row)]
        if not numbered:
            continue
        sheet_names.append(sheet_name)
        # Trim trailing empty columns.
        width = max(
            max((i + 1 for i, cell in enumerate(row) if cell), default=0)
            for _, row in numbered
        )
        numbered = [(n, row[:width]) for n, row in numbered]
        header_number, header_row = numbered[0]
        header = [
            f"[工作表 {sheet_name}]",
            f"[R{header_number}] " + " | ".join(header_row),
        ]
        lines = [f"[R{n}] " + " | ".join(row) for n, row in numbered[1:]]
        table = {
            "index": len(sheet_names),
            "sheet": sheet_name,
            "rows": [row for _, row in numbered],
            "row_numbers": [n for n, _ in numbered],
        }
        sheet_blocks = _blocks(lines, header) or ["\n".join(header)]
        for position, text in enumerate(sheet_blocks):
            pages.append(ParsedPage(
                page=len(pages) + 1,
                text=text,
                tables=[table] if position == 0 else [],
            ))

    return ParsedDocument(
        document_id=path.stem,
        source_type="spreadsheet",
        source_name=path.name,
        source_path=str(path),
        sha256=calculate_sha256(path),
        pages=pages,
        metadata={
            "page_count": len(pages),
            "page_unit": "sheet_block",
            "sheets": sheet_names,
            "text_char_count": sum(len(page.text) for page in pages),
            "empty_page_count": 0,
            "needs_ocr": False,
        },
    )

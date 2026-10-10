"""Human spot check of the historical backfill: 20 random documents (10 related-party, 10 guarantee), at most 5 records
each, written to data/eval/hist_spot_check.xlsx for a teammate to mark against the PDF.

The sample is drawn with a fixed seed BEFORE anyone looks at the outputs, from documents with at least one record.
Fields checked: counterparty, amount, unit, page (and, for related-party rows, last year's actual amount).
Omissions are counted per document (records in the PDF that the system did not extract).

    python scripts/make_spot_check.py
"""
from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
FREEZE = ROOT / "outputs" / "hist_freeze"
SOURCES = ROOT / "data" / "manifests" / "hist_sources.csv"
OUT = ROOT / "data" / "eval" / "hist_spot_check.xlsx"
SEED, PER_TASK, PER_DOC = 20261009, 10, 5
REL = {"parent_or_controlling_shareholder": "控股股东/母公司", "sister_company": "同一集团兄弟公司",
       "wholly_owned_subsidiary": "全资子公司", "controlled_subsidiary": "控股子公司", "associate": "联营/合营",
       "other_related_party": "其他关联方"}


def records(task: str, doc: dict) -> list[dict]:
    if task == "related_party":
        return [{"对方": t.get("counterparty"), "关系": REL.get(t.get("relationship"), t.get("relationship")),
                 "类别": t.get("category_text") or t.get("transaction_category"), "金额": t.get("estimated_amount"),
                 "单位": t.get("estimated_unit"), "上年实际": t.get("prior_year_actual_amount"),
                 "页码": t.get("source_page"), "系统给的原文": t.get("evidence_text")} for t in doc.get("transactions") or []]
    return [{"对方": e.get("guaranteed_party"), "关系": REL.get(e.get("relationship"), e.get("relationship")),
             "类别": e.get("event_type"), "金额": e.get("guarantee_amount"), "单位": e.get("guarantee_unit"),
             "上年实际": "", "页码": e.get("source_page"), "系统给的原文": e.get("evidence_text")}
            for e in doc.get("events") or []]


def main() -> None:
    with SOURCES.open(encoding="utf-8-sig", newline="") as f:
        meta = {r["file_name"].rsplit(".", 1)[0]: r for r in csv.DictReader(f)}
    rng = random.Random(SEED)
    picked = []
    for task in ("related_party", "guarantee"):
        pool = []
        for path in sorted((FREEZE / task).glob("*.json")):
            doc = json.loads(path.read_text(encoding="utf-8"))
            if records(task, doc) and path.stem in meta:
                pool.append((task, path.stem, doc))
        picked += rng.sample(pool, PER_TASK)
    wb = Workbook()
    guide = wb.active
    guide.title = "说明"
    for line in ["怎么核对（约 3–4 小时）",
                 "目的：检查系统从公告里抽出的数字和名称对不对、有没有漏。标准答案只有一个：公告原文。",
                 "",
                 "一、逐条核对（表“逐条核对”，62 行）",
                 "1. 点“公告链接”打开原文 PDF，翻到“页码”那一页（页码按 PDF 第几页算，不是印在页脚的页码）。",
                 "2. 在那一页的表格里找到“对方”这一行，然后逐项判断，在黄色格子里选：对 / 错 / 原文没有。",
                 "   对方对吗：公司名称与原文一致即为对（简称与全称指同一家公司也算对）。",
                 "   金额对吗：关联交易看“预计金额（本年度预计）”那一列；担保看“担保额度 / 担保金额”。数字完全相同为对，千分位逗号不算差别。",
                 "   单位对吗：看表头或表格上方写的单位（元 / 万元 / 亿元），与“单位”列一致为对。",
                 "   页码对吗：这条记录确实在该页为对；在别的页为错，在备注写正确页码。",
                 "   上年实际对吗（只有关联交易有）：看“上年实际发生金额”那一列。注意有的表有三列数字（本年预计、本年截至披露日已发生、上年实际），要看表头选对列。",
                 "3. 选“错”时，在备注写原文的正确值，例如“原文为 632,065.26”。原文里根本找不到这条记录时，各项都选“原文没有”。",
                 "",
                 "二、漏抽统计（表“漏抽统计”，20 行）",
                 "4. 每份公告数一下原文主表（关联交易预计表或担保明细表）里一共有多少条记录，填在“原文表格中的记录条数”。不数“合计”“小计”行。",
                 "   系统抽出的条数已经写在旁边；两者的差就是漏抽（或多抽），由我们来算，你只需要数原文。",
                 "",
                 "三、交回",
                 "5. 拿不准的写进备注，不要猜。全部填完把这个 Excel 发回即可；不要改动白色格子。"]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 140
    ws = wb.create_sheet("逐条核对")
    head = ["编号", "公告", "公司", "公告日期", "公告标题", "公告链接", "记录", "对方", "关系", "类别", "金额", "单位", "上年实际",
            "页码", "系统给的原文", "对方对吗", "金额对吗", "单位对吗", "页码对吗", "上年实际对吗", "备注"]
    ws.append(head)
    misses = wb.create_sheet("漏抽统计")
    misses.append(["编号", "公告", "公司", "公告标题", "公告链接", "系统抽出条数", "原文表格中的记录条数（你来数）", "备注"])
    choice = DataValidation(type="list", formula1='"对,错,原文没有"', allow_blank=True)
    ws.add_data_validation(choice)
    for i, (task, stem, doc) in enumerate(picked, 1):
        m = meta[stem]
        recs = records(task, doc)
        for j, r in enumerate(recs[:PER_DOC], 1):
            ws.append([f"D{i:02d}", stem, m["security_name"], m["announcement_date"], m["announcement_title"], m["source_url"],
                       f"{j}/{len(recs)}", r["对方"], r["关系"], r["类别"], r["金额"], r["单位"], r["上年实际"], r["页码"],
                       (r["系统给的原文"] or "")[:200], "", "", "", "", "" if task == "related_party" else "—", ""])
        misses.append([f"D{i:02d}", stem, m["security_name"], m["announcement_title"], m["source_url"], len(recs), "", ""])
    choice.add(f"P2:T{ws.max_row}")
    fill = PatternFill("solid", fgColor="FFF4CC")
    for row in ws.iter_rows(min_row=2, min_col=16, max_col=21):
        for c in row:
            c.fill = fill
    for row in misses.iter_rows(min_row=2, min_col=7, max_col=8):
        for c in row:
            c.fill = fill
    for sheet in (ws, misses):
        for c in sheet[1]:
            c.font = Font(bold=True)
        sheet.freeze_panes = "B2"
    for col, width in zip("ABCDEFGHIJKLMNOPQRSTU", [6, 30, 10, 11, 36, 20, 6, 26, 14, 12, 10, 6, 10, 6, 40, 9, 9, 9, 9, 11, 24]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=15, max_col=15):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{len(picked)} documents, {ws.max_row - 1} records -> {OUT.relative_to(ROOT)} (seed {SEED})")


if __name__ == "__main__":
    sys.exit(main())

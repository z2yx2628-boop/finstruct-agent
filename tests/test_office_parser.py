import zipfile

from openpyxl import Workbook

from src.document_parser import parse_document
from src.office_parser import parse_docx, parse_spreadsheet

DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:r><w:t>湖南华菱钢铁股份有限公司关于为子公司提供担保的公告</w:t></w:r></w:p>
<w:p><w:r><w:t>公司为子公司提供担保，</w:t></w:r><w:r><w:t>担保方式为连带责任保证。</w:t></w:r></w:p>
<w:tbl>
<w:tr><w:tc><w:p><w:r><w:t>被担保方</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>担保额度（万元）</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>湖南华菱涟源钢铁有限公司</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>300,000</w:t></w:r></w:p></w:tc></w:tr>
</w:tbl>
<w:p><w:r><w:t>公司无逾期担保。</w:t></w:r></w:p>
</w:body>
</w:document>"""


def make_docx(path):
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", DOCUMENT_XML)


def test_docx_paragraphs_and_table(tmp_path):
    path = tmp_path / "notice.docx"
    make_docx(path)

    document = parse_docx(path)
    text = document.pages[0].text

    assert document.source_type == "word"
    assert "[P1] 湖南华菱钢铁股份有限公司关于为子公司提供担保的公告" in text
    assert "[P2] 公司为子公司提供担保，担保方式为连带责任保证。" in text
    assert "湖南华菱涟源钢铁有限公司 | 300,000" in text
    assert "[P3] 公司无逾期担保。" in text
    assert ["湖南华菱涟源钢铁有限公司", "300,000"] in document.pages[0].tables[0]["rows"]


def test_xlsx_rows_are_numbered_and_header_repeats(tmp_path):
    path = tmp_path / "detail.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "担保明细"
    sheet.append(["被担保方", "担保额度（万元）", None])
    for index in range(1, 121):
        sheet.append([f"子公司{index:03d}有限公司", 1000 + index, None])
    workbook.create_sheet("空表")
    workbook.save(path)

    document = parse_spreadsheet(path)

    assert document.source_type == "spreadsheet"
    assert document.metadata["sheets"] == ["担保明细"]
    assert len(document.pages) > 1
    for page in document.pages:
        assert page.text.startswith("[工作表 担保明细]\n[R1] 被担保方 | 担保额度（万元）")
    assert "[R2] 子公司001有限公司 | 1001" in document.pages[0].text
    all_text = "\n".join(page.text for page in document.pages)
    assert "[R121] 子公司120有限公司 | 1120" in all_text
    assert document.pages[0].tables[0]["row_numbers"][:2] == [1, 2]


def test_gbk_csv_is_decoded(tmp_path):
    path = tmp_path / "detail.csv"
    path.write_bytes("被担保方,金额\n南京钢铁股份有限公司,5\n".encode("gb18030"))

    document = parse_document(path)

    assert document.source_type == "spreadsheet"
    assert "[R2] 南京钢铁股份有限公司 | 5" in document.pages[0].text


def test_office_files_are_routed(tmp_path):
    path = tmp_path / "notice.DOCX"
    make_docx(path)

    assert parse_document(path).source_type == "word"


FIXTURES = __import__("pathlib").Path(__file__).parent / "fixtures"


def test_legacy_doc_paragraphs_and_table():
    document = parse_document(FIXTURES / "legacy_notice.doc")
    text = document.pages[0].text

    assert document.source_type == "word"
    assert document.metadata["legacy_format"] is True
    assert "[P1] 湖南华菱钢铁股份有限公司关于为子公司提供担保的公告" in text
    assert "[P2] 公司拟为湖南华菱涟源钢铁有限公司提供担保，担保金额为30亿元。" in text
    assert "华菱涟钢 | 300,000" in text
    assert ["华菱涟钢", "300,000"] in document.pages[0].tables[0]["rows"]


def test_legacy_xls_values_formulas_and_long_string_table():
    document = parse_document(FIXTURES / "legacy_detail.xls")

    assert document.source_type == "spreadsheet"
    assert document.metadata["sheets"] == ["担保明细", "大表"]
    first = document.pages[0].text
    assert first.startswith("[工作表 担保明细]\n[R1] 被担保方 | 银行 | 担保金额（万元） | 持股比例")
    assert "[R2] 湖南华菱涟源钢铁有限公司 | 中国银行股份有限公司湘潭分行 | 300000 | 1" in first
    assert "[R3] 华菱钢铁（香港）国际贸易有限公司 | 交通银行 | 50000.5 | 0.8591" in first
    assert "[R4] 合计 |  | 350000.5" in first  # cached formula result
    # 700 distinct strings force the shared-string table across CONTINUE records.
    big = "\n".join(page.text for page in document.pages[1:])
    assert "[R701] 700 | 测试钢铁子公司第700号有限责任公司" in big
    assert big.count("有限责任公司") == 700


def test_unsupported_legacy_file_gives_clear_error(tmp_path):
    import pytest
    from src.legacy_office import LegacyOfficeError

    fake = tmp_path / "old.doc"
    fake.write_bytes(b"not an ole file")
    with pytest.raises(LegacyOfficeError):
        parse_docx(fake)
